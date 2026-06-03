"""Экспорт статистики в Excel с картинками и данными книг."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from io import BytesIO

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image as XLImage
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

import db


def export_statistics_to_excel(output_path: str | Path) -> None:
    """Экспортирует статистику чтения в Excel файл с картинками и данными."""
    if not HAS_OPENPYXL:
        raise ImportError("Требуется пакет 'openpyxl'. Установите его: pip install openpyxl pillow")
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Статистика чтения"
    
    # Стили
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )
    
    # Ширина колонок
    ws.column_dimensions['A'].width = 25  # Название
    ws.column_dimensions['B'].width = 18  # Автор
    ws.column_dimensions['C'].width = 15  # Жанр
    ws.column_dimensions['D'].width = 12  # Страниц
    ws.column_dimensions['E'].width = 15  # Начало
    ws.column_dimensions['F'].width = 15  # Конец
    ws.column_dimensions['G'].width = 18  # Статус
    ws.column_dimensions['H'].width = 15  # Формат
    ws.column_dimensions['I'].width = 20  # Теги
    ws.column_dimensions['J'].width = 60  # Обложка (для изображения)
    
    # Высота для изображений
    ws.row_dimensions[1].height = 25
    
    # Заголовки
    headers = ["Название", "Автор", "Жанр", "Страниц", "Начало чтения", "Конец чтения", 
               "Статус", "Формат", "Теги", "Обложка"]
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    
    # Данные книг
    books = db.list_all_books()
    
    for row_idx, book in enumerate(books, 2):
        ws.row_dimensions[row_idx].height = 80  # Высота для изображения
        
        title = book.get("title") or ""
        author = book.get("author") or ""
        genre = book.get("genre") or ""
        page_count = book.get("page_count") or 0
        
        # Просто показываем количество страниц
        pages_text = str(page_count)
        
        date_started = book.get("date_started")
        if hasattr(date_started, 'isoformat'):
            date_started = date_started.isoformat()[:10]
        else:
            date_started = str(date_started or "")[:10]
        
        date_finished = book.get("date_finished")
        if hasattr(date_finished, 'isoformat'):
            date_finished = date_finished.isoformat()[:10]
        else:
            date_finished = str(date_finished or "")[:10]
        
        status = book.get("reading_status") or "planned"
        status_display = {"planned": "В планах", "reading": "Читаю", "finished": "Прочитано"}.get(status, status)
        
        format_type = book.get("format_type") or "paper"
        format_display = {"paper": "Бумажная", "ebook": "Электронная", "audiobook": "Аудиокнига"}.get(format_type, format_type)
        
        # Получаем теги
        try:
            tags = db.get_book_tags(book["id"])
            tags_text = ", ".join(tags)
        except:
            tags_text = ""
        
        # Заполняем ячейки
        cells_data = [
            title,
            author,
            genre,
            pages_text,
            date_started,
            date_finished,
            status_display,
            format_display,
            tags_text,
            ""  # Место для обложки
        ]
        
        for col_idx, value in enumerate(cells_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = value
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = border
        
        # Добавляем изображение если есть
        cover_path = book.get("cover_path")
        if cover_path and isinstance(cover_path, str) and Path(cover_path).is_file():
            try:
                img = XLImage(cover_path)
                img.width = 70
                img.height = 100
                ws.add_image(img, f"J{row_idx}")
            except Exception:
                # Если не удалось добавить изображение, просто пропускаем
                pass
    
    # Добавляем лист со статистикой
    ws_stats = wb.create_sheet("Общая статистика")
    ws_stats.column_dimensions['A'].width = 30
    ws_stats.column_dimensions['B'].width = 20
    
    # Статистика
    all_books = db.list_all_books()
    finished_books = [b for b in all_books if b.get("reading_status") == "finished"]
    reading_books = [b for b in all_books if b.get("reading_status") == "reading"]
    planned_books = [b for b in all_books if b.get("reading_status") == "planned"]
    
    total_pages_read = sum(int(b.get("page_count") or 0) for b in finished_books)
    
    stats_data = [
        ["Статистика", ""],
        ["Всего книг", len(all_books)],
        ["Прочитано", len(finished_books)],
        ["В процессе", len(reading_books)],
        ["В планах", len(planned_books)],
        ["Всего страниц прочитано", total_pages_read],
    ]
    
    for row_idx, (label, value) in enumerate(stats_data, 1):
        cell_a = ws_stats.cell(row=row_idx, column=1)
        cell_b = ws_stats.cell(row=row_idx, column=2)
        cell_a.value = label
        cell_b.value = value
        
        if row_idx == 1:
            cell_a.font = Font(bold=True, size=12)
        else:
            cell_a.border = border
            cell_b.border = border
    
    wb.save(output_path)


def get_export_status() -> tuple[bool, str]:
    """Проверяет доступность Excel экспорта и возвращает статус."""
    if HAS_OPENPYXL:
        return True, "Excel экспорт доступен"
    else:
        return False, "Требуется установить: pip install openpyxl pillow"
